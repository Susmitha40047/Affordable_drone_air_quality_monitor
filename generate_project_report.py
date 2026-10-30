from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "Affordable_Drone_Air_Quality_Monitoring_System_Report.docx"
PDF_PATH = ROOT / "Affordable_Drone_Air_Quality_Monitoring_System_Report.pdf"

TITLE = "Affordable Drone Air Quality Monitoring System"
TEAM_NUMBER = "5"
COURSE = "A Project Based Learning Report"
SUBTITLE = (
    "Submitted in partial fulfilment of the requirements for the award of the degree of "
    "Bachelor of Technology"
)
DEPARTMENT = "The Department of Electronics and Communication Engineering"
COURSE_CODE = "EMBEDDED PROTOTYPE (23 SDEC14A)"
INSTITUTION = "Koneru Lakshmaiah Education Foundation, Aziz Nagar"
LOCATION = "Aziz Nagar - 500075"
YEAR = "APRIL - 2026"
GUIDE = "Guide: ____________________"

TEAM_MEMBERS = [
    "2310040047 : G. S. Susmitha",
    "2310040081 : S. Vaishnavi",
    "2310040061 : P. Shruthi",
    "2310040034 : L. Archana",
]

SECTIONS = [
    (
        "Declaration",
        [
            "We hereby declare that the project titled Affordable Drone Air Quality Monitoring System is the original work carried out by us under the academic guidance of the department for the purpose of Project Based Learning. The work presented in this report has not been submitted elsewhere for the award of any degree or diploma. The content included in this report is prepared for academic review and project evaluation.",
        ],
    ),
    (
        "Abstract",
        [
            "This project presents a software-based implementation of an affordable drone air quality monitoring system. The proposed work demonstrates how a drone-inspired platform can be used to monitor environmental pollution at different locations through a complete software simulation. Conventional air quality monitoring stations are expensive and fixed in specific locations, which limits their coverage. The proposed system addresses this limitation by simulating a mobile drone platform that travels across multiple GPS points and collects air-quality information dynamically.",
            "The developed system generates readings for PM2.5, PM10, CO2, gas index, temperature, humidity, altitude, and speed. These values are processed by a Flask-based backend and displayed through a real-time dashboard containing live metrics, route visualization, pollution trend charts, and alert messages. The generated mission data is also stored in CSV format for later analysis and reporting. The software prototype provides a complete working model of a drone-based environmental monitoring system and serves as a practical foundation for future integration with real drone hardware, embedded sensors, wireless communication, and cloud platforms.",
        ],
    ),
    (
        "Introduction",
        [
            "Air pollution has become one of the most serious environmental challenges affecting public health and sustainable living. Pollution monitoring is essential in urban locations, industrial zones, educational campuses, and public gathering spaces where pollutant concentration can change from place to place. Traditional air quality monitoring stations offer accurate measurement, but they are expensive, fixed in location, and difficult to deploy in large numbers.",
            "A drone-based air quality monitoring system offers a mobile and flexible solution. A drone can travel across different locations, collect pollution readings, and provide a better understanding of spatial variation in air quality. Because the present project had a very short development timeline, the complete concept was implemented as a software prototype. This approach allows demonstration of the overall system architecture, data flow, dashboard interface, and decision logic without waiting for hardware procurement, drone integration, or sensor calibration.",
        ],
    ),
    (
        "Problem Statement",
        [
            "Existing fixed air quality monitoring systems are costly and cannot monitor multiple locations quickly. There is a need for an affordable and portable monitoring model that can simulate or collect data from different points and present the information in real time for analysis and decision making.",
        ],
    ),
    (
        "Need for the Project",
        [
            "Air pollution varies from one location to another, especially in urban and semi-urban zones. A single fixed monitoring station cannot represent all local conditions. A mobile system such as a drone can help collect data from several points within a short duration. Since hardware implementation requires time for integration, calibration, and testing, a software prototype is useful to validate the concept, system flow, and dashboard behavior before building the actual hardware version.",
        ],
    ),
    (
        "Objectives",
        [
            "To design a complete software prototype of an affordable drone-based air quality monitoring system.",
            "To simulate drone movement using predefined GPS mission points.",
            "To generate realistic air quality values such as PM2.5, PM10, CO2, gas index, temperature, and humidity.",
            "To display live mission data using a web dashboard with charts and route mapping.",
            "To generate alerts when the simulated air pollution level crosses threshold conditions.",
            "To store mission data in CSV format for later observation and analysis.",
        ],
    ),
    (
        "Literature Review",
        [
            "Air quality monitoring has been an active area of research due to rising environmental concerns and its impact on human health. Traditional systems mainly relied on fixed monitoring stations equipped with expensive analytical instruments. These installations are suitable for detailed observation but are often limited to government or industrial deployments because of their cost and infrastructure requirements.",
            "Recent developments in embedded systems and low-cost sensors have made compact monitoring systems more practical. Microcontroller platforms such as ESP32 and Arduino are widely used for environmental monitoring because they support sensor integration, local processing, and wireless communication. Researchers have also explored Internet of Things based monitoring, where environmental data is transmitted to cloud dashboards for remote supervision and historical analysis.",
            "Mobile monitoring has gained importance because pollution is not uniformly distributed across an environment. Drones and autonomous mobile platforms are increasingly studied for environmental sensing because they can travel across large areas and collect location-specific information. Existing literature shows that combining environmental sensing, geographic tagging, and real-time visualization can significantly improve the usefulness of air quality monitoring systems. The current project follows this direction and presents a complete software model of such a system for academic demonstration.",
        ],
    ),
    (
        "Existing System",
        [
            "Existing air quality monitoring solutions are usually fixed installations. They provide good accuracy, but they are expensive, cover limited area, and require higher maintenance. Fixed systems also do not support easy spatial monitoring across multiple nearby points. For student-level implementation, these systems are difficult to replicate due to hardware cost and setup complexity.",
        ],
    ),
    (
        "Proposed System",
        [
            "The proposed system is a software simulation of a drone carrying air quality sensors. The drone follows a virtual mission route represented by GPS coordinates. At every route point, the simulator generates environmental readings and sends them to a backend application. The backend updates a web dashboard in real time and logs the readings into a CSV file.",
            "The software is divided into a mission simulator, sensor data generator, Flask backend, and dashboard user interface. The dashboard displays live metric cards, route mapping, trend charts, and alerts. This architecture demonstrates how a real hardware system would behave when connected to actual sensors and telemetry modules.",
        ],
    ),
    (
        "Advantages over Existing System",
        [
            "The proposed system is affordable, scalable, and suitable for academic demonstration. It supports multi-point monitoring, real-time visualization, and alert generation. It also serves as a design framework for future hardware deployment using ESP32, gas sensors, GPS modules, and a drone flight controller.",
        ],
    ),
    (
        "System Architecture",
        [
            "The complete system architecture consists of four major blocks: drone route simulation, air quality data generation, backend processing, and dashboard visualization. The route simulation module defines how the virtual drone moves through GPS points. The sensor generator creates realistic air-quality values. The backend provides APIs for the web interface and stores all mission records. The dashboard presents the mission status, charts, maps, and alerts to the user.",
            "Architecture Flow: Drone Route Simulator -> Sensor Data Generator -> Flask Backend -> Dashboard and CSV Logger.",
        ],
    ),
    (
        "Software Requirements",
        [
            "Programming Language: Python",
            "Backend Framework: Flask",
            "Frontend Technologies: HTML, CSS, JavaScript",
            "Charts and Visualization: Chart.js",
            "Map Integration: Leaflet.js",
            "Data Storage: CSV",
            "Development Environment: Visual Studio Code or equivalent Python editor",
        ],
    ),
    (
        "Implementation",
        [
            "The implementation begins with a Python-based simulator that generates a circular route around a reference geographic location. The simulator produces one mission reading at regular time intervals. Each record contains latitude, longitude, altitude, speed, PM2.5, PM10, CO2, gas index, temperature, humidity, air-quality label, and alert status.",
            "The backend is implemented using Flask. It exposes APIs for current mission state, historical data, and mission controls such as start, stop, and reset. The frontend dashboard uses these APIs to fetch data periodically and update the user interface. The dashboard includes mission control buttons, metric cards, line charts, a live route map, and an alert panel that highlights unsafe readings. Every generated record is written into a CSV file to simulate data logging in a real air quality monitoring mission.",
        ],
    ),
    (
        "Module Description",
        [
            "Module 1 - Mission Simulator: Generates drone route coordinates and mission sequence.",
            "Module 2 - Sensor Data Generator: Produces simulated pollution and environmental readings.",
            "Module 3 - Backend API: Handles start, stop, reset, and data communication.",
            "Module 4 - Dashboard UI: Displays metrics, route map, charts, and alerts.",
            "Module 5 - CSV Logger: Stores mission results for later analysis.",
        ],
    ),
    (
        "Working Procedure",
        [
            "1. The user opens the dashboard in the browser.",
            "2. The user starts the mission using the Start Mission button.",
            "3. The simulator begins generating route points and air-quality values.",
            "4. The backend sends the generated data to the dashboard through APIs.",
            "5. The map displays the drone position and route path.",
            "6. Charts update the pollution and environmental trends in real time.",
            "7. Alerts appear when PM2.5 or CO2 exceeds a threshold.",
            "8. All readings are saved into a CSV file for future analysis.",
        ],
    ),
    (
        "Testing and Validation",
        [
            "The project was validated by checking simulator execution, dashboard updates, chart rendering, route plotting, alert logic, and CSV generation. The application was tested by running the mission repeatedly and observing whether generated values were reflected correctly in the user interface. This confirmed the consistency of backend processing and frontend visualization.",
        ],
    ),
    (
        "Results and Output",
        [
            "The developed prototype successfully demonstrates the expected workflow of a drone-based air quality monitoring system. The mission dashboard displays live drone position, simulated sensor values, pollution charts, and alert notifications. The route map shows how the drone covers multiple geographic locations. The logged CSV file contains the full mission dataset, which can be used for result analysis or report preparation.",
            "The simulation clearly shows how a real drone-based monitoring system can collect and visualize environmental data across multiple points. Even without hardware implementation, the software model provides a convincing demonstration of system behavior and user interaction.",
        ],
    ),
    (
        "Advantages",
        [
            "Affordable academic prototype for quick implementation.",
            "Demonstrates the full project workflow without hardware delay.",
            "Supports multi-location monitoring concept through a mobile drone model.",
            "Provides real-time visualization and alert generation.",
            "Can be upgraded later into a real hardware-integrated system.",
        ],
    ),
    (
        "Applications",
        [
            "Urban pollution monitoring",
            "Industrial emission observation",
            "Traffic area analysis",
            "Environmental research projects",
            "Smart campus demonstrations",
        ],
    ),
    (
        "Limitations",
        [
            "The present version is a software simulation and not a hardware deployment.",
            "Actual sensor calibration, flight stability, and wireless transmission are not included.",
            "The generated data is simulated for demonstration purposes.",
        ],
    ),
    (
        "Future Scope",
        [
            "Integrate ESP32 and real air quality sensors.",
            "Connect the platform to a real drone and flight controller.",
            "Add live telemetry and cloud data storage.",
            "Build a mobile application for remote monitoring.",
            "Apply machine learning to predict pollution trends.",
        ],
    ),
    (
        "Conclusion",
        [
            "The project successfully demonstrates a complete software-based model of an affordable drone air quality monitoring system. It simulates drone movement, air-quality sensing, live visualization, alert generation, and mission logging in a single integrated application. The developed prototype provides a practical and presentable solution for academic review and shows the feasibility of using drones for environmental monitoring.",
            "Although the current work is implemented as a simulation, it forms a strong base for future expansion into a real-time embedded and drone-enabled system. The project satisfies the core objective of demonstrating affordable, location-aware air quality monitoring using a mobile platform concept.",
        ],
    ),
    (
        "References",
        [
            "Flask Official Documentation.",
            "Chart.js Documentation.",
            "Leaflet.js Documentation.",
            "Research concepts related to air quality monitoring, embedded systems, and drone-based sensing.",
        ],
    ),
]


def set_page_border(section) -> None:
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "18")
        element.set(qn("w:color"), "808080")
        pg_borders.append(element)
    sect_pr.append(pg_borders)


def add_docx_title_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    set_page_border(section)

    for _ in range(3):
        document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(18)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COURSE)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(16)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SUBTITLE)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DEPARTMENT)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(13)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COURSE_CODE)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(GUIDE)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Submitted by\nTeam Number: {TEAM_NUMBER}")
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(12)

    for member in TEAM_MEMBERS:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(member)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(INSTITUTION)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(12)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(LOCATION)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(YEAR)
    run.font.name = "Times New Roman"
    run.bold = True
    run.font.size = Pt(12)


def add_docx_body(document: Document) -> None:
    document.add_section(WD_SECTION.NEW_PAGE)
    for section in document.sections[1:]:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        set_page_border(section)

    for title, paragraphs in SECTIONS:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = "Times New Roman"
        run.bold = True
        run.font.size = Pt(14)

        for text in paragraphs:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Inches(0.3)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)


def build_docx() -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    add_docx_title_page(document)
    add_docx_body(document)
    document.save(DOCX_PATH)


def build_pdf() -> None:
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontName="Times-Bold",
        fontSize=18,
        alignment=TA_CENTER,
        leading=24,
        spaceAfter=12,
    )
    center_style = ParagraphStyle(
        "CenterStyle",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=12,
        alignment=TA_CENTER,
        leading=16,
        spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=14,
        textColor=colors.black,
        spaceBefore=10,
        spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=12,
        leading=17,
        alignment=TA_JUSTIFY,
        firstLineIndent=18,
        spaceAfter=8,
    )

    story = []
    story.append(Spacer(1, 1.25 * inch))
    story.append(Paragraph(TITLE, title_style))
    story.append(Paragraph(COURSE, center_style))
    story.append(Paragraph(SUBTITLE, center_style))
    story.append(Paragraph(DEPARTMENT, center_style))
    story.append(Paragraph(COURSE_CODE, center_style))
    story.append(Paragraph(GUIDE, center_style))
    story.append(Spacer(1, 0.25 * inch))
    story.append(Paragraph(f"Submitted by<br/>Team Number: {TEAM_NUMBER}", center_style))
    for member in TEAM_MEMBERS:
        story.append(Paragraph(member, center_style))
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(INSTITUTION, center_style))
    story.append(Paragraph(LOCATION, center_style))
    story.append(Paragraph(YEAR, center_style))
    story.append(PageBreak())

    for title, paragraphs in SECTIONS:
        story.append(Paragraph(title, heading_style))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph.replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 0.08 * inch))

    doc.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
